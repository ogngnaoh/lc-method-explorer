"""Audit source workbooks without modifying them or fitting models."""
from collections import Counter, defaultdict
import hashlib
import json
import math
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def clean(value):
    return ' '.join(value.split()) if isinstance(value, str) else value


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def main():
    manifest = json.loads((ROOT / 'data/source/mcmrt_manifest.json').read_text())
    methods, issues = [], []
    compounds = defaultdict(set)
    structures = defaultdict(set)
    observations = []
    for item in manifest['files']:
        path = ROOT / 'data/raw/mcmrt' / item['name']
        assert hashlib.sha256(path.read_bytes()).hexdigest() == item['sha256'], path
        method_id = path.stem.split()[-1]
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        assert book.sheetnames == ['RT', 'LC setups'], (path, book.sheetnames)
        rows = list(book['RT'].values)
        headers = [clean(v) for v in rows[0]]
        assert headers[:11] == ['MCMRT Number', 'Compound Name', 'IUPAC Name', 'Formula',
                               'CAS Number', 'Pubchem Number', 'Isomeric SMILES', 'InChI',
                               'Retention Factor (k)', 'RT (min)', 'RSD'], (path, headers)
        assert not any(r[0] is None and any(v is not None for v in r[:11])
                       for r in rows[1:]), (path, 'Nonempty row without compound ID')
        records = [dict(zip(headers[:11], r[:11])) for r in rows[1:] if r[0] is not None]
        if any(any(v is not None for v in r[11:]) for r in rows):
            issues.append({'method': method_id, 'issue': 'nonempty extra columns'})
        setup_rows = list(book['LC setups'].values)
        marker = next(i for i, r in enumerate(setup_rows) if clean(r[0]) == 'Gradient elution program')
        setup = {clean(r[0]): clean(r[1]) for r in setup_rows[:marker]}
        gradient = [list(r[1:4]) for r in setup_rows[marker+1:] if any(v is not None for v in r[1:4])]
        assert all(len(g) == 3 and all(number(v) for v in g) for g in gradient), path
        assert all(g[1] > 0 and 0 <= g[2] <= 100 for g in gradient), path
        assert all(a[0] < b[0] for a, b in zip(gradient, gradient[1:])), path
        ids = [r['MCMRT Number'] for r in records]
        assert len(ids) == len(set(ids)), path
        valid = 0
        for r in records:
            cid = r['MCMRT Number']
            structures[cid].add((clean(r['Isomeric SMILES']), clean(r['InChI'])))
            rt = r['RT (min)']
            if not number(rt) or rt <= 0:
                issues.append({'method': method_id, 'compound': cid, 'issue': 'missing or invalid RT', 'value': rt})
                continue
            valid += 1
            compounds[cid].add(method_id)
            observations.append({'method': method_id, 'compound': cid, 'rt_min': rt})
        methods.append({'method_id': method_id, 'source_file': path.name,
                        'rows_with_compound_id': len(records), 'measured_rt': valid,
                        'settings': setup, 'gradient_time_flow_percent_b': gradient})
        book.close()
    stats = {
        'files': len(methods), 'rows_with_compound_id': sum(m['rows_with_compound_id'] for m in methods),
        'measured_rt': len(observations), 'unique_source_compound_ids': len(structures),
        'compounds_measured_in_all_methods': sum(len(v) == len(methods) for v in compounds.values()),
        'coverage_histogram_methods_per_compound': dict(sorted(Counter(len(v) for v in compounds.values()).items())),
        'inconsistent_structures_by_source_id': {str(k): sorted(v) for k, v in structures.items() if len(v) > 1},
        'distinct_column_names': sorted({m['settings']['Analytical column'] for m in methods}),
        'distinct_mobile_phase_pairs': sorted({(m['settings']['Mobile phase A'], m['settings']['Mobile phase B']) for m in methods}),
        'distinct_time_percent_b_programs': len({tuple((g[0],g[2]) for g in m['gradient_time_flow_percent_b']) for m in methods}),
        'column_temperatures_c': sorted({m['settings']['Column temperature (°C)'] for m in methods}),
        'run_times_min': sorted({m['gradient_time_flow_percent_b'][-1][0] for m in methods}),
        'rt_min_range': [min(r['rt_min'] for r in observations), max(r['rt_min'] for r in observations)],
        'issues': issues, 'methods': methods,
        'limitations': ['Compound counts use source identifiers, not RDKit-standardized structures.',
                       'No model training or evaluation has been performed.']}
    output = ROOT / 'reports'
    output.mkdir(exist_ok=True)
    (output / 'data_audit.json').write_text(json.dumps(stats, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps({k:v for k,v in stats.items() if k != 'methods'}, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
